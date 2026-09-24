#!/usr/bin/env python3
"""Deterministically account for every Tropical Pacific workbook creature.

Only rows meeting the documented, conservative inclusion policy enter the app pack.
Everything else remains visible in the outcome and review reports. Workbook media is
never read because it is reference-only and is not licensed for the application.
"""
from __future__ import annotations

import argparse, csv, json, re, sys
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'CatalogReview'))
from catalog_review import apply_decisions, update_manifest

BASELINE_SIZE = 40
BATCH_SIZE = 100
COLORS = {'black','blue','brown','gray','green','olive','orange','red','silver','white','yellow'}
MARKINGS = {'barbels','beak','eye stripe','fin edge','patches','saddles','shell','spines','spots','stripes','tail','teeth'}
HABITATS = {'anemone','deep','lagoon','mangrove','open water','reef','rubble','sand','seagrass','shallow','surface','wall','wreck'}
BEHAVIORS = {'burrowing','cleaning','feeding','grazing','hiding','hovering','resting','schooling','solitary','swimming'}

# Workbook group names are source classifications, not app vocabulary.  Keep this
# allow-list explicit: suffix matching would incorrectly classify jellyfish,
# cuttlefish, starfish, and any future OCR corruption as fish.
FISH_GROUPS = {
    'Anemonefishes', 'Angelfishes', 'Anthias', 'Barracudas', 'Batfishes',
    'Bigeyes', 'Blennies', 'Boxfishes', 'Brotulas', 'Butterflyfishes',
    'Cardinalfishes', 'Chubs', 'Coral Breams', 'Cornetfishes', 'Damselfishes',
    'Dartfishes', 'Devilfishes', 'Diamondfishes', 'Dolphinfishes', 'Dottybacks',
    'Dragonets', 'Eel-tailed Catfishes', 'Emperors', 'Filefishes', 'Fish',
    'Fishes', 'Flagtails', 'Flashlightfishes', 'Flatheads', 'Flounders',
    'Frogfishes', 'Fusiliers', 'Goatfishes', 'Gobies', 'Groupers', 'Grunters',
    'Gurnards', 'Hawkfishes', 'Jawfishes', 'Jacks', 'Lionfishes',
    'Lizardfishes', 'Milkfish', 'Mojarras', 'Moorish Idol', 'Mullets',
    'Needlefishes', 'Parrotfishes', 'Pearl Fishes', 'Pearl Perches',
    'Pineconefishes', 'Pipefishes', 'Ponyfishes', 'Porcupinefishes', 'Puffers',
    'Rabbitfishes', 'Sand Divers', 'Sandperches', 'Scorpionfishes',
    'Scorpionfishes/Lionfishes', 'Sea Basses/Anthias', 'Sea Moths',
    'Seabasses/Basslets', 'Shrimpfishes', 'Signalfishes', 'Silversides',
    'Snappers', 'Snooks', 'Soapfishes', 'Soles', 'Spadefishes',
    'Squirrelfishes', 'Stargazers', 'Stonefishes', 'Surgeonfishes',
    'Sweetlips', 'Sweepers', 'Tarpons', 'Threadfins', 'Tilefishes',
    'Triplefins', 'Trumpetfishes', 'Triggerfishes', 'Tunas & Mackerels',
    'Waspfishes', 'Wormfishes', 'Wrasses',
}
EEL_GROUPS = {'Conger Eels', 'Garden Eels', 'Morays', 'Snake Eels'}
SEAHORSE_GROUPS = {'Seahorses'}
SHARK_GROUPS = {'Bamboo Sharks', 'Cat Sharks', 'Nurse Sharks', 'Requiem Sharks', 'Sharks', 'Whale Sharks', 'Wobbegongs'}
RAY_GROUPS = {'Guitarfishes', 'Manta Rays', 'Mantas', 'Rays', 'Stingrays', 'Wedgefishes'}
GROUP_CATEGORY_MAP = {
    **{name.casefold(): 'fish' for name in FISH_GROUPS},
    **{name.casefold(): 'eel' for name in EEL_GROUPS},
    **{name.casefold(): 'seahorse' for name in SEAHORSE_GROUPS},
    **{name.casefold(): 'shark' for name in SHARK_GROUPS},
    **{name.casefold(): 'ray' for name in RAY_GROUPS},
}

def rows(sheet):
    iterator=sheet.iter_rows(values_only=True); header=[str(value) for value in next(iterator)]
    return [dict(zip(header,row)) for row in iterator if any(value is not None for value in row)]
def clean(value): return re.sub(r'\s+',' ',str(value or '')).strip()
def number(value): return float(value) if value is not None and clean(value) else None
def terms(text, vocabulary):
    low=text.lower()
    return [value for value in sorted(vocabulary) if any(re.search(r'(?<![a-z])'+re.escape(candidate)+r'(?![a-z])',low) for candidate in {value,value.rstrip('s'),value+'s'})]
def iso_date(value):
    if value is None: return None
    result=value.isoformat() if hasattr(value,'isoformat') else str(value)
    return result+'T00:00:00Z' if re.fullmatch(r'\d{4}-\d{2}-\d{2}',result) else result.replace('+00:00','Z')
def normalized(value): return clean(value).casefold()

def canonical_category(value):
    """Resolve an exact, normalized workbook group or return a review diagnostic."""
    source=clean(value); canonical=GROUP_CATEGORY_MAP.get(source.casefold())
    if canonical: return canonical, ''
    if not source: return None, 'missing workbook category'
    return None, f'unrecognized workbook category: {source}'

def review_reasons(creature, source_traits, source):
    reasons=[]
    if clean(creature.get('occurrence_status'))!='presence_supported': reasons.append('regional presence is not source-supported')
    if clean(creature.get('identity_confidence'))!='high' or number(creature.get('identity_confidence_score')) != 100: reasons.append('identity confidence is below the inclusion threshold')
    if clean(creature.get('text_transcription_confidence'))!='high': reasons.append('account transcription confidence is not high')
    if not clean(creature.get('common_name')) or not clean(creature.get('scientific_name_printed')): reasons.append('printed identity is incomplete')
    if not source_traits: reasons.append('no identification description')
    elif not any(clean(trait.get('transcription_confidence'))=='high' for trait in source_traits): reasons.append('identification descriptions are not high-confidence transcriptions')
    if source is None: reasons.append('species-account source locator is missing')
    elif not clean(source.get('source_id')) or not clean(source.get('source_locator')): reasons.append('species-account source ID or locator is missing')
    return reasons

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--workbook',default='DiveID_Tropical_Pacific_Internal_Consistency_Cleaned.xlsx')
    parser.add_argument('--output',default='DiveID/Resources/IdentificationPacks/TropicalPacific')
    parser.add_argument('--report',default='Reports/TropicalPacificImportReport.json')
    parser.add_argument('--outcomes',default='Reports/TropicalPacificOutcomes.csv')
    parser.add_argument('--review-queue',default='Reports/TropicalPacificReviewQueue.csv')
    parser.add_argument('--review-decisions',default='Data/CatalogReview/ReviewDecisions.json')
    args=parser.parse_args()
    workbook=load_workbook(args.workbook,data_only=True,read_only=True)
    creatures=rows(workbook['Creatures']); traits=defaultdict(list); sources={}
    for trait in rows(workbook['Traits']): traits[clean(trait.get('creature_id'))].append(trait)
    for source in rows(workbook['Sources']):
        creature_id=clean(source.get('creature_id'))
        if creature_id and creature_id not in sources: sources[creature_id]=source

    assessed=[]
    for index, creature in enumerate(creatures,1):
        creature_id=clean(creature.get('creature_id'))
        descriptions=[trait for trait in traits[creature_id] if clean(trait.get('trait_type')) in ('identification_description','variation_description')]
        assessed.append((index,creature,descriptions,sources.get(creature_id),review_reasons(creature,descriptions,sources.get(creature_id))))
    candidates=[item for item in assessed if not item[4]]
    common_counts=Counter(normalized(item[1].get('common_name')) for item in candidates)
    scientific_counts=Counter(normalized(item[1].get('scientific_name_printed')) for item in candidates)
    for item in candidates:
        if common_counts[normalized(item[1].get('common_name'))]>1: item[4].append('common name duplicates another otherwise eligible row')
        if scientific_counts[normalized(item[1].get('scientific_name_printed'))]>1: item[4].append('printed scientific name duplicates another otherwise eligible row')

    eligible=sorted((item for item in assessed if not item[4]),key=lambda item:(int(item[1].get('source_book_page') or 99999),clean(item[1].get('creature_id'))))
    baseline_ids=[clean(item[1].get('creature_id')) for item in eligible[:BASELINE_SIZE]]
    profiles=[]; outcomes=[]
    for index, creature, descriptions, source, reasons in assessed:
        creature_id=clean(creature.get('creature_id'))
        if not creature_id or not re.fullmatch(r'[0-9a-fA-F-]{36}',creature_id):
            outcome='excluded'; reasons=['missing or malformed creature_id']+reasons
        elif reasons: outcome='pending_review'
        else: outcome='included'
        evidence=' | '.join(clean(trait.get('value')) for trait in descriptions if clean(trait.get('value')))
        source_category=clean(creature.get('category')); canonical, category_diagnostic=canonical_category(source_category)
        outcomes.append({'workbook_row':index+1,'creature_id':creature_id,'common_name':clean(creature.get('common_name')),'scientific_name_printed':clean(creature.get('scientific_name_printed')),'outcome':outcome,'reasons':'; '.join(dict.fromkeys(reasons)),'workbook_human_review_required':clean(creature.get('human_review_required')),'source_id':clean(creature.get('source_id')),'source_book_page':clean(creature.get('source_book_page')),'source_pdf_page':clean(creature.get('source_pdf_page')),'source_tile':clean(creature.get('source_tile')),'source_category':source_category,'canonical_category':canonical or '','category_diagnostic':category_diagnostic,'identity_confidence':clean(creature.get('identity_confidence')),'identity_confidence_score':clean(creature.get('identity_confidence_score')),'text_transcription_confidence':clean(creature.get('text_transcription_confidence')),'source_title_common_ocr':clean(creature.get('source_title_common_ocr')),'source_title_scientific_ocr':clean(creature.get('source_title_scientific_ocr')),'identification_evidence':evidence})
        if outcome!='included': continue
        high_traits=[clean(trait.get('value')) for trait in descriptions if clean(trait.get('value')) and clean(trait.get('transcription_confidence'))=='high']
        trait_text=' '.join(high_traits); max_size=number(creature.get('max_size_cm'))
        profiles.append({'id':creature_id,'commonName':clean(creature.get('common_name')),'scientificName':clean(creature.get('scientific_name_printed')),'aliases':[],
          'categories':[canonical] if canonical else [],'colors':terms(trait_text,COLORS),'markings':terms(trait_text,MARKINGS),'bodyShapes':[],
          'habitats':terms(trait_text,HABITATS),'regions':[],'behaviors':terms(trait_text,BEHAVIORS),'keywords':[],
          'minimumSizeCentimeters':None,'maximumSizeCentimeters':max_size,'minimumDepthMeters':number(creature.get('depth_min_m')),'maximumDepthMeters':number(creature.get('depth_max_m')),
          'summary':high_traits[0],'distinguishingFeatures':high_traits,'typicalHabitat':'','geographicRange':clean(creature.get('range_detail_raw')),
          'cautions':[],'imageAssetName':None,'regionalOccurrence':'unknown','regionalOccurrenceNotes':clean(creature.get('occurrence_basis')) or None,'subregions':[],
          'appearanceVariants':[],'similarSpecies':[],'bundledImage':None,
          'dataSources':[{'stableSourceID':clean(source.get('source_id')),'sourceName':clean(source.get('citation')),'sourceURL':clean(source.get('species_specific_url')),'citationReference':clean(source.get('source_locator')),'reviewedFields':[field.strip() for field in clean(source.get('supported_fields')).split(';') if field.strip()],'accessedDate':iso_date(source.get('accessed_date')),'sourceLicense':clean(source.get('licence')) or None}],
          'review':{'status':'draft','reviewerNotes':'Scientific name is the source-printed identity, not an independently accepted taxonomy. Source-page and taxonomy review remain required.','reviewDate':None,'verifiedBy':None},
          'taxonomy':None,'measurements':None,'tailShape':None,'mouthAndHeadShape':[],'finAndSpineClues':[]})

    profiles.sort(key=lambda profile:(profile['commonName'].casefold(),profile['id']))
    output=Path(args.output); output.mkdir(parents=True,exist_ok=True)
    manifest={'id':'tropical-pacific','schemaVersion':1,'packVersion':2,'displayName':'Tropical Pacific (Experimental)','shortDescription':'Source-traceable draft Tropical Pacific catalogue; not verified for publication','geographicScope':'Workbook accounts with source-supported Tropical Pacific presence; structural inclusion does not verify identity, biology, or abundance.','regionAliases':['Tropical Pacific','Pacific','Indo-Pacific','Fiji','Hawaii','Australia','Philippines','Indonesia'],'speciesCount':len(profiles),'speciesResourceName':'Creatures','imageSubdirectory':'Images','includedWithApp':True,'lastDataReviewDate':None,'includedRecordCount':len(profiles),'humanReviewedRecordCount':0,'publicationEligibleRecordCount':0}
    decision_path=Path(args.review_decisions)
    decision_outcomes=apply_decisions(profiles,json.loads(decision_path.read_text(encoding='utf-8')),manifest['id']) if decision_path.is_file() else []
    update_manifest(manifest,profiles)
    (output/'Creatures.json').write_text(json.dumps(profiles,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    (output/'PackManifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    fields=list(outcomes[0]); Path(args.outcomes).parent.mkdir(parents=True,exist_ok=True)
    with Path(args.outcomes).open('w',newline='',encoding='utf-8') as handle:
        writer=csv.DictWriter(handle,fieldnames=fields,lineterminator='\n'); writer.writeheader(); writer.writerows(outcomes)
    # Included rows remain draft and are also queued when the workbook asks for review.
    pending=[row for row in outcomes if row['outcome']=='pending_review' or row['workbook_human_review_required'].lower() in ('true','1','yes')]
    with Path(args.review_queue).open('w',newline='',encoding='utf-8') as handle:
        writer=csv.DictWriter(handle,fieldnames=fields,lineterminator='\n'); writer.writeheader(); writer.writerows(pending)
    counts=Counter(row['outcome'] for row in outcomes)
    source_category_counts=Counter(row['source_category'] or '(missing)' for row in outcomes if row['outcome']=='included')
    canonical_category_counts=Counter(row['canonical_category'] or '(unresolved)' for row in outcomes if row['outcome']=='included')
    unresolved_category_counts=Counter(row['source_category'] or '(missing)' for row in outcomes if row['category_diagnostic'])
    included_unresolved_counts=Counter(row['source_category'] or '(missing)' for row in outcomes if row['outcome']=='included' and row['category_diagnostic'])
    report={'workbook':args.workbook,'workbookRowCount':len(creatures),'recordCounts':{'structurallyIncluded':len(profiles),'humanReviewed':manifest['humanReviewedRecordCount'],'publicationEligible':manifest['publicationEligibleRecordCount'],'reviewQueue':len(pending)},'reviewDecisionOutcomes':decision_outcomes,'outcomeCounts':{'included':counts['included'],'pendingReview':counts['pending_review'],'excluded':counts['excluded']},'publicationPolicy':'Structural import, parsing, normalization, and search testing never confer publication approval. Promotion requires traceable source evidence plus named, dated human review with notes.','categoryNormalization':{'method':'casefolded, collapsed-whitespace exact allow-list; no substring matching','mapping':dict(sorted((key,value) for key,value in GROUP_CATEGORY_MAP.items())),'includedSourceCategoryCounts':dict(sorted(source_category_counts.items())),'includedCanonicalCategoryCounts':dict(sorted(canonical_category_counts.items())),'includedUnresolvedCategoryCounts':dict(sorted(included_unresolved_counts.items())),'unresolvedWorkbookCategoryCounts':dict(sorted(unresolved_category_counts.items()))},'selectionPolicy':{'rules':['source-supported presence','identity confidence high with score 100','high account transcription confidence','complete printed identity','high-confidence identification description','source ID and locator','unique app identity'],'baselineSize':BASELINE_SIZE,'baselineIDs':baseline_ids,'batchSize':BATCH_SIZE,'batches':[{'number':n//BATCH_SIZE+1,'start':n+1,'end':min(n+BATCH_SIZE,len(profiles)),'recordIDs':[p['id'] for p in profiles[n:n+BATCH_SIZE]]} for n in range(0,len(profiles),BATCH_SIZE)]},'bundledRecordIDs':[p['id'] for p in profiles],'mediaImported':0,'mediaPolicy':'Workbook media remains reference_only_not_licensed_for_app and is never read or imported.','schemaResolution':'Unknown abundance is encoded as regionalOccurrence=unknown. Printed scientific names remain source identity; taxonomy and measurements are null rather than asserting an accepted name or measurement type.','determinism':'Workbook order is used only for row reporting; bundle, baseline, reasons, and batches use explicit stable ordering. Reports contain no run timestamp.'}
    Path(args.report).write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print(f"processed {len(creatures)} rows: bundled {counts['included']}, pending {counts['pending_review']}, excluded {counts['excluded']}")
if __name__=='__main__': main()
